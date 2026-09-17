"""Tests for xlsx_inspection.py. The real Anthropic client is always mocked
here - these tests never make a network call or consume API credit.
"""

import re
import sys
import tempfile
import types
import unittest
import uuid
import zipfile
from pathlib import Path
from unittest.mock import MagicMock, patch

sys.path.insert(0, str(Path(__file__).parent.parent))

import anthropic
import httpx2
import openpyxl

import documents
import store
import xlsx_inspection


def strip_dimension_declarations(path: Path) -> None:
    """Rewrites an .xlsx in place with every sheet's <dimension .../>
    element removed - reproduces a real-world workbook (confirmed live
    against an actual deal's financial model) saved by a tool that omits
    or mis-states that declaration, which makes openpyxl's read_only mode
    report max_row/max_column as None for every sheet instead of computing
    them."""
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = Path(tmp.name)
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp_path, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                data = re.sub(rb"<dimension[^/]*/>", b"", data)
            zout.writestr(item, data)
    path.write_bytes(tmp_path.read_bytes())
    tmp_path.unlink()

FAKE_SECRET = "sk-ant-api03-XLSX-INSPECTION-TEST-FAKE-SECRET-DO-NOT-LEAK"


def make_status_error(cls, status_code: int, error_type: str, message: str):
    request = httpx2.Request("POST", "https://api.anthropic.com/v1/files")
    response = httpx2.Response(
        status_code, request=request, json={"error": {"type": error_type, "message": message}}
    )
    return cls(message, response=response, body=response.json())


def build_sample_workbook(path: Path, marker: str = "") -> bytes:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    assert ws1 is not None
    ws1.title = "Revenue Model"
    ws1["B7"] = 100
    ws1["B8"] = "=B7*1.1"
    ws1.merge_cells("D1:E1")
    ws1["D1"] = "Merged header"
    ws1["C10"] = "='[ExternalBook.xlsx]Sheet1'!A1"  # external-link-style formula
    ws2 = wb.create_sheet("Assumptions")
    ws2.sheet_state = "hidden"
    ws2["A1"] = f"hidden data {marker}"
    wb.save(path)
    return path.read_bytes()


def fake_text_block(text: str):
    return types.SimpleNamespace(type="text", text=text)


def fake_tool_use_block(name: str, tool_input: dict):
    return types.SimpleNamespace(type="server_tool_use", name=name, input=tool_input)


def fake_bash_result_block(stdout: str = "", stderr: str = "", return_code: int = 0):
    content = types.SimpleNamespace(
        type="bash_code_execution_result", stdout=stdout, stderr=stderr, return_code=return_code
    )
    return types.SimpleNamespace(type="bash_code_execution_tool_result", content=content)


def fake_bash_error_block(error_code: str = "execution_time_exceeded"):
    content = types.SimpleNamespace(type="bash_code_execution_tool_result_error", error_code=error_code)
    return types.SimpleNamespace(type="bash_code_execution_tool_result", content=content)


def fake_response(
    content_blocks,
    model: str = "claude-opus-5",
    stop_reason: str = "end_turn",
    in_tok: int = 1000,
    out_tok: int = 500,
    container_id: str | None = "container_1",
    code_execution_requests: int = 1,
    stop_details=None,
):
    usage = types.SimpleNamespace(
        input_tokens=in_tok,
        output_tokens=out_tok,
        server_tool_use=types.SimpleNamespace(code_execution_requests=code_execution_requests),
    )
    container = types.SimpleNamespace(id=container_id) if container_id else None
    return types.SimpleNamespace(
        content=content_blocks,
        model=model,
        usage=usage,
        stop_reason=stop_reason,
        stop_details=stop_details,
        container=container,
    )


def _capturing_upload(captured: dict):
    """A files.upload side_effect that reads the file handle while it's
    still open (inside the caller's `with open(...)` block) and stashes
    the bytes for the test to inspect afterward - the handle is closed by
    the time inspect_workbook() returns, so reading it later would fail.
    """

    def upload(*, file, **kwargs):
        filename, fileobj, mime_type = file
        captured["filename"] = filename
        captured["bytes"] = fileobj.read()
        captured["mime_type"] = mime_type
        return types.SimpleNamespace(id="file_abc123")

    return upload


def _stream_cm(response):
    mock_stream = MagicMock()
    mock_stream.get_final_message.return_value = response
    cm = MagicMock()
    cm.__enter__ = MagicMock(return_value=mock_stream)
    cm.__exit__ = MagicMock(return_value=False)
    return cm


class XlsxInspectionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._tmpdir = tempfile.TemporaryDirectory()
        cls._original_data_dir = documents.DATA_DIR
        documents.DATA_DIR = Path(cls._tmpdir.name) / "DealLabData"
        cls._schema = f"test_{uuid.uuid4().hex}"
        cls._original_schema = store.SCHEMA
        store.ensure_schema(cls._schema)
        store.SCHEMA = cls._schema
        store.init_db()
        documents.init_documents_db()
        cls.project = store.create_project("Project Falcon", "test project for workbook inspection")

    @classmethod
    def tearDownClass(cls):
        store.SCHEMA = cls._original_schema
        store.drop_schema(cls._schema)
        documents.DATA_DIR = cls._original_data_dir
        cls._tmpdir.cleanup()

    def setUp(self):
        self._env_patcher = patch.dict(
            "os.environ", {"ANTHROPIC_API_KEY": FAKE_SECRET, "ANTHROPIC_MODEL": "claude-opus-5"}
        )
        self._env_patcher.start()

        scratch = Path(self._tmpdir_scratch()) / "sample.xlsx"
        self.workbook_bytes = build_sample_workbook(scratch, marker=self.id())
        result = documents.save_uploaded_file(self.project.id, "model.xlsx", "", self.workbook_bytes)
        assert result.status == "success" and result.document is not None
        self.document = result.document

    def _tmpdir_scratch(self) -> str:
        if not hasattr(self, "_scratch_dir"):
            self._scratch_dir_obj = tempfile.TemporaryDirectory()
            self._scratch_dir = self._scratch_dir_obj.name
        return self._scratch_dir

    def tearDown(self):
        self._env_patcher.stop()
        if hasattr(self, "_scratch_dir_obj"):
            self._scratch_dir_obj.cleanup()

    def _mock_client(self, stream_responses=None, stream_side_effect=None, upload_side_effect=None, delete_side_effect=None):
        mock_client = MagicMock()
        if upload_side_effect is not None:
            mock_client.files.upload.side_effect = upload_side_effect
        else:
            mock_client.files.upload.return_value = types.SimpleNamespace(id="file_abc123")

        if delete_side_effect is not None:
            mock_client.files.delete.side_effect = delete_side_effect
        else:
            mock_client.files.delete.return_value = types.SimpleNamespace(id="file_abc123", type="file_deleted")

        if stream_side_effect is not None:
            mock_client.messages.stream.side_effect = stream_side_effect
        elif stream_responses is not None:
            mock_client.messages.stream.side_effect = [_stream_cm(r) for r in stream_responses]

        return patch("xlsx_inspection.anthropic.Anthropic", return_value=mock_client), mock_client

    # -- basic error paths -------------------------------------------------

    def test_missing_api_key(self):
        with patch.dict("os.environ", {"ANTHROPIC_API_KEY": ""}):
            outcome = xlsx_inspection.inspect_workbook(self.document)
        self.assertFalse(outcome.success)
        self.assertFalse(outcome.transmitted)
        self.assertEqual(outcome.error_type, "missing_api_key")

    def test_not_a_workbook_rejected_before_transmission(self):
        txt_result = documents.save_uploaded_file(self.project.id, "notes.txt", "", b"plain text")
        assert txt_result.document is not None
        outcome = xlsx_inspection.inspect_workbook(txt_result.document)
        self.assertFalse(outcome.success)
        self.assertFalse(outcome.transmitted)
        self.assertEqual(outcome.error_type, "not_workbook")

    def test_oversized_workbook_rejected_before_transmission(self):
        with patch.object(xlsx_inspection, "MAX_WORKBOOK_SOURCE_BYTES", len(self.workbook_bytes) - 1):
            patcher, mock_client = self._mock_client()
            with patcher:
                outcome = xlsx_inspection.inspect_workbook(self.document)
        self.assertFalse(outcome.success)
        self.assertFalse(outcome.transmitted)
        self.assertEqual(outcome.error_type, "oversized_workbook")
        mock_client.files.upload.assert_not_called()

    def test_missing_stored_file(self):
        result = documents.save_uploaded_file(self.project.id, "vanishing.xlsx", "", self.workbook_bytes + b"x")
        assert result.document is not None
        documents.stored_file_path(result.document).unlink()
        outcome = xlsx_inspection.inspect_workbook(result.document)
        self.assertFalse(outcome.success)
        self.assertFalse(outcome.transmitted)
        self.assertEqual(outcome.error_type, "missing_file")

    # -- upload / transmission ----------------------------------------------

    def test_uploads_exact_original_bytes(self):
        response = fake_response([fake_text_block("## Workbook Identification and Purpose\nA model.")])
        captured: dict = {}
        patcher, mock_client = self._mock_client(
            stream_responses=[response], upload_side_effect=_capturing_upload(captured)
        )
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(self.document)

        self.assertTrue(outcome.success)
        self.assertEqual(captured["bytes"], self.workbook_bytes)
        self.assertEqual(captured["mime_type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_no_other_project_document_is_uploaded(self):
        other_bytes = build_sample_workbook(Path(self._tmpdir_scratch()) / "other.xlsx", marker="other")
        other_result = documents.save_uploaded_file(self.project.id, "other.xlsx", "", other_bytes)
        assert other_result.status == "success"

        response = fake_response([fake_text_block("## Workbook Identification and Purpose\nA model.")])
        captured: dict = {}
        patcher, mock_client = self._mock_client(
            stream_responses=[response], upload_side_effect=_capturing_upload(captured)
        )
        with patcher:
            xlsx_inspection.inspect_workbook(self.document)

        self.assertEqual(mock_client.files.upload.call_count, 1)
        self.assertEqual(captured["bytes"], self.workbook_bytes)
        self.assertNotEqual(captured["bytes"], other_bytes)

    def test_upload_failure_reports_error_without_calling_messages(self):
        error = make_status_error(anthropic.APIStatusError, 413, "invalid_request_error", "file too large")
        patcher, mock_client = self._mock_client(upload_side_effect=error)
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(self.document)
        self.assertFalse(outcome.success)
        self.assertEqual(outcome.error_type, "oversized_workbook")
        mock_client.messages.stream.assert_not_called()

    # -- cleanup -------------------------------------------------------------

    def test_remote_cleanup_attempted_and_recorded_on_success(self):
        response = fake_response([fake_text_block("## Workbook Identification and Purpose\nA model.")])
        patcher, mock_client = self._mock_client(stream_responses=[response])
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(self.document)

        self.assertTrue(outcome.remote_cleanup_attempted)
        self.assertTrue(outcome.remote_cleanup_succeeded)
        mock_client.files.delete.assert_called_once_with("file_abc123")

    def test_remote_cleanup_failure_is_recorded_not_hidden(self):
        response = fake_response([fake_text_block("## Workbook Identification and Purpose\nA model.")])
        patcher, mock_client = self._mock_client(
            stream_responses=[response], delete_side_effect=RuntimeError("boom")
        )
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(self.document)

        self.assertTrue(outcome.success)  # analysis itself still succeeded
        self.assertTrue(outcome.remote_cleanup_attempted)
        self.assertFalse(outcome.remote_cleanup_succeeded)

    # -- citation parsing and existence verification -------------------------

    def test_citation_existence_verification(self):
        text = (
            "Revenue is 100 ('Revenue Model'!B7 [value]) from a formula "
            "('Revenue Model'!B8 [formula]). A bogus sheet citation "
            "('Nonexistent'!Z9 [value]) and an out-of-range cell "
            "('Revenue Model'!B999999 [value])."
        )
        response = fake_response([fake_text_block(text)])
        patcher, _ = self._mock_client(stream_responses=[response])
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(self.document)

        self.assertTrue(outcome.success)
        self.assertTrue(outcome.verification_available)
        citations = [s.citation for s in outcome.segments if s.type == "citation"]
        self.assertEqual(len(citations), 4)
        self.assertTrue(citations[0].exists)
        self.assertTrue(citations[1].exists)
        self.assertFalse(citations[2].exists)  # nonexistent sheet
        self.assertFalse(citations[3].exists)  # out of range

    def test_citation_with_trailing_value_text_is_still_recognized(self):
        # Regression test: confirmed live, Claude reliably uses the
        # sheet!ref/[kind] shape but sometimes appends the cell's actual
        # value or label before the closing paren, e.g.
        # ('Assumptions'!B4 [value] 1000000) - a stricter format that
        # required the closing paren to immediately follow "[kind]" missed
        # more than half of a real response's citations entirely.
        text = (
            "Base revenue is set ('Revenue Model'!B7 [value] 100) and the "
            "header reads ('Revenue Model'!D1 [label] \"Merged header\"), with "
            "a formula computing growth "
            "('Revenue Model'!B8 [formula] `=B7*1.1`)."
        )
        response = fake_response([fake_text_block(text)])
        patcher, _ = self._mock_client(stream_responses=[response])
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(self.document)

        self.assertTrue(outcome.success)
        citations = [s.citation for s in outcome.segments if s.type == "citation"]
        self.assertEqual(len(citations), 3)
        self.assertEqual([c.kind for c in citations], ["value", "label", "formula"])
        self.assertTrue(all(c.exists for c in citations))

    def test_citation_with_nested_parens_in_formula_text_is_parsed_cleanly(self):
        # Regression test: a formula citation like NPV(...) contains its
        # own parentheses, so a naive "match up to the next )" approach
        # stops at the formula's inner closing paren and leaves a stray
        # ")" dangling as plain text afterward. Confirmed live against a
        # real NPV formula citation.
        text = (
            "The total is computed here "
            "('Revenue Model'!B8 [formula] `=NPV(0.1,B7,B7)`) and nothing "
            "else follows."
        )
        response = fake_response([fake_text_block(text)])
        patcher, _ = self._mock_client(stream_responses=[response])
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(self.document)

        self.assertTrue(outcome.success)
        segments = outcome.segments
        self.assertEqual(len(segments), 3)
        self.assertEqual(segments[0].type, "text")
        self.assertEqual(segments[1].type, "citation")
        self.assertEqual(segments[1].citation.raw_text, "('Revenue Model'!B8 [formula] `=NPV(0.1,B7,B7)`)")
        self.assertTrue(segments[1].citation.exists)
        self.assertEqual(segments[2].type, "text")
        self.assertEqual(segments[2].text, " and nothing else follows.")
        # No stray ")" leaking into the trailing text segment.
        self.assertNotIn(")", segments[2].text)

    def test_hidden_sheet_and_merged_cell_and_external_link_do_not_break_verification(self):
        text = (
            "Hidden assumptions are tracked ('Assumptions'!A1 [label]). "
            "The merged header is present ('Revenue Model'!D1 [label]). "
            "An external link formula exists ('Revenue Model'!C10 [formula])."
        )
        response = fake_response([fake_text_block(text)])
        patcher, _ = self._mock_client(stream_responses=[response])
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(self.document)

        self.assertTrue(outcome.success)
        self.assertTrue(outcome.verification_available)
        citations = [s.citation for s in outcome.segments if s.type == "citation"]
        self.assertTrue(all(c.exists for c in citations))

    def test_workbook_missing_dimension_declaration_does_not_false_reject_citations(self):
        # Regression test: confirmed live against a real deal's financial
        # model. When a workbook lacks a valid <dimension> element,
        # openpyxl's fast read_only mode can't compute sheet bounds and
        # returns None for every sheet - naively treating that as "1 row, 1
        # column" rejected every real citation into the file as out of
        # range, even though the citations were correct and the sheet
        # names matched exactly.
        scratch = Path(self._tmpdir_scratch()) / "no_dimension.xlsx"
        workbook_bytes = build_sample_workbook(scratch, marker=self.id())
        strip_dimension_declarations(scratch)
        workbook_bytes = scratch.read_bytes()
        result = documents.save_uploaded_file(self.project.id, "no_dimension.xlsx", "", workbook_bytes)
        assert result.document is not None

        text = "Revenue is 100 ('Revenue Model'!B7 [value]) from a formula ('Revenue Model'!B8 [formula])."
        response = fake_response([fake_text_block(text)])
        patcher, _ = self._mock_client(stream_responses=[response])
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(result.document)

        self.assertTrue(outcome.success)
        self.assertTrue(outcome.verification_available)
        citations = [s.citation for s in outcome.segments if s.type == "citation"]
        self.assertEqual(len(citations), 2)
        self.assertTrue(all(c.exists for c in citations))

    def test_missing_citations_are_empty_not_invented(self):
        response = fake_response([fake_text_block("## Summary\nA general uncited remark with no cell reference.")])
        patcher, _ = self._mock_client(stream_responses=[response])
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(self.document)

        self.assertTrue(outcome.success)
        citation_segments = [s for s in outcome.segments if s.type == "citation"]
        self.assertEqual(citation_segments, [])

    def test_malformed_local_copy_marks_verification_unavailable_not_a_failure(self):
        # The workbook Claude was actually sent might still be a genuine
        # xlsx, but if our own on-disk copy can't be opened for structural
        # verification, that must not be reported as Claude's analysis
        # failing - only that verification could not be performed.
        result = documents.save_uploaded_file(self.project.id, "corrupt.xlsx", "", b"not a real xlsx file")
        assert result.document is not None

        response = fake_response([fake_text_block("('Sheet1'!A1 [value]) some finding.")])
        patcher, _ = self._mock_client(stream_responses=[response])
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(result.document)

        self.assertTrue(outcome.success)
        self.assertFalse(outcome.verification_available)
        self.assertEqual(outcome.verification_unavailable_reason, "malformed_workbook")
        citation = [s.citation for s in outcome.segments if s.type == "citation"][0]
        self.assertIsNone(citation.exists)

    def test_encrypted_local_copy_marks_verification_unavailable(self):
        encrypted_bytes = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64
        result = documents.save_uploaded_file(self.project.id, "encrypted.xlsx", "", encrypted_bytes)
        assert result.document is not None

        response = fake_response([fake_text_block("No citations here.")])
        patcher, _ = self._mock_client(stream_responses=[response])
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(result.document)

        self.assertTrue(outcome.success)
        self.assertFalse(outcome.verification_available)
        self.assertEqual(outcome.verification_unavailable_reason, "encrypted_workbook")

    # -- pause_turn / truncation / tool errors -------------------------------

    def test_pause_turn_continues_and_succeeds(self):
        paused = fake_response([fake_text_block("Partial thought...")], stop_reason="pause_turn")
        finished = fake_response([fake_text_block("## Workbook Identification and Purpose\nDone.")])
        patcher, mock_client = self._mock_client(stream_responses=[paused, finished])
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(self.document)

        self.assertTrue(outcome.success)
        self.assertEqual(mock_client.messages.stream.call_count, 2)
        # The continuation call must reuse the same container.
        _, second_kwargs = mock_client.messages.stream.call_args_list[1]
        self.assertEqual(second_kwargs["container"], "container_1")

    def test_still_paused_after_max_continuations_is_an_error(self):
        paused = fake_response([fake_text_block("Still working...")], stop_reason="pause_turn")
        patcher, mock_client = self._mock_client(
            stream_responses=[paused] * (xlsx_inspection.MAX_PAUSE_CONTINUATIONS + 1)
        )
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(self.document)

        self.assertFalse(outcome.success)
        self.assertEqual(outcome.error_type, "still_paused")
        self.assertTrue(outcome.remote_cleanup_attempted)

    def test_max_tokens_truncation_is_a_failure_with_partial_output_preserved(self):
        response = fake_response(
            [fake_text_block("## Workbook Identification and Purpose\nPartial")], stop_reason="max_tokens"
        )
        patcher, _ = self._mock_client(stream_responses=[response])
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(self.document)

        self.assertFalse(outcome.success)
        self.assertEqual(outcome.error_type, "truncated_response")
        self.assertIsNotNone(outcome.segments)

    def test_tool_call_trace_is_captured_including_errors(self):
        blocks = [
            fake_tool_use_block("bash_code_execution", {"command": "python inspect.py"}),
            fake_bash_result_block(stdout="Revenue Model, Assumptions", return_code=0),
            fake_tool_use_block("bash_code_execution", {"command": "slow_command"}),
            fake_bash_error_block("execution_time_exceeded"),
            fake_text_block("## Workbook Identification and Purpose\nDone."),
        ]
        response = fake_response(blocks)
        patcher, _ = self._mock_client(stream_responses=[response])
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(self.document)

        self.assertTrue(outcome.success)
        kinds = [entry["kind"] for entry in outcome.tool_trace]
        self.assertEqual(kinds, ["tool_use", "bash_result", "tool_use", "bash_error"])
        self.assertEqual(outcome.tool_trace[3]["error_code"], "execution_time_exceeded")
        # Regression test: confirmed live, the installed SDK's
        # response.usage.server_tool_use never exposes a
        # code_execution_requests field (only web_fetch/web_search
        # counts), even when real tool calls happened - so this is derived
        # from our own trace instead of that always-empty SDK field.
        self.assertEqual(outcome.usage["code_execution_requests"], 2)

    def test_tool_trace_stdout_is_capped(self):
        long_stdout = "x" * (xlsx_inspection.TOOL_TRACE_TEXT_CAP + 500)
        blocks = [fake_bash_result_block(stdout=long_stdout), fake_text_block("Done ('Revenue Model'!B7 [value]).")]
        response = fake_response(blocks)
        patcher, _ = self._mock_client(stream_responses=[response])
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(self.document)

        captured_stdout = outcome.tool_trace[0]["stdout"]
        self.assertLess(len(captured_stdout), len(long_stdout))
        self.assertIn("truncated", captured_stdout)

    # -- generic API error classification ------------------------------------

    def test_invalid_api_key(self):
        error = make_status_error(anthropic.AuthenticationError, 401, "authentication_error", "invalid x-api-key")
        patcher, _ = self._mock_client(stream_side_effect=error)
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(self.document)
        self.assertFalse(outcome.success)
        self.assertEqual(outcome.error_type, "invalid_api_key")

    def test_insufficient_credit_real_world_400_shape(self):
        # Regression test: this exact status/type/message combination was
        # observed live against a genuinely out-of-credit account, and was
        # originally misclassified as a generic "unexpected_error" because
        # the check only recognized a dedicated 402/"billing_error" shape.
        error = make_status_error(
            anthropic.APIStatusError,
            400,
            "invalid_request_error",
            "Your credit balance is too low to access the Anthropic API. "
            "Please go to Plans & Billing to upgrade or purchase credits.",
        )
        patcher, mock_client = self._mock_client(stream_side_effect=error)
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(self.document)
        self.assertFalse(outcome.success)
        self.assertEqual(outcome.error_type, "insufficient_credit")
        # A billing failure still deletes the uploaded file.
        mock_client.files.delete.assert_called_once()

    def test_secret_never_appears_in_outcome(self):
        error = make_status_error(
            anthropic.AuthenticationError, 401, "authentication_error", f"invalid api key: {FAKE_SECRET}"
        )
        patcher, _ = self._mock_client(stream_side_effect=error)
        with patcher:
            outcome = xlsx_inspection.inspect_workbook(self.document)
        self.assertNotIn(FAKE_SECRET, str(outcome.error_message))


if __name__ == "__main__":
    unittest.main()
