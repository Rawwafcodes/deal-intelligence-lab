"""Export builders for the Deal Workspace (Milestone 9): native XLSX
findings/requests registers (via openpyxl, already a dependency - no new
one added for this) and printer-friendly HTML for the executive memo and
the combined decision package, which a user saves as a PDF through the
browser's own print dialog rather than a server-side PDF library.

Every export is built only from data already loaded by the caller (the
workspace's own findings/requests/memo plus the analysis and project
records) - never from the filesystem, the Anthropic API key, or any other
secret. See DISCLAIMER and the identifying header every export carries.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from html import escape
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

import cross_format_analyses
import store
import workspaces

DISCLAIMER = (
    "This export reflects findings from an AI-assisted document reconciliation together with a human "
    "reviewer's own judgments and workflow state. It is a working diligence artifact, not investment, "
    "legal, accounting, or tax advice, and does not represent an automated or final transaction decision. "
    "Every finding requires independent professional verification before being relied upon."
)

_FINDINGS_COLUMNS = [
    ("Finding ID", "id"),
    ("Origin", "origin_label"),
    ("Title", "title"),
    ("Classification", "classification"),
    ("AI severity", "ai_severity"),
    ("Human-adjusted severity", "adjusted_severity"),
    ("Effective severity", "effective_severity"),
    ("Review status", "review_status"),
    ("Resolution status", "resolution_status"),
    ("Assigned owner", "assigned_owner"),
    ("Duplicate?", "duplicate_label"),
    ("Duplicate of", "duplicate_of"),
    ("Explanation", "explanation"),
    ("PDF evidence", "pdf_evidence_text"),
    ("Workbook evidence", "workbook_evidence_text"),
    ("Commercial/financial relevance", "commercial_relevance"),
    ("Uncertainty", "uncertainty"),
    ("Recommended action", "recommended_action"),
    ("Management response", "management_response"),
    ("Reviewer notes", "reviewer_notes"),
    ("Due date", "due_date_text"),
    ("Last updated", "updated_at"),
]

_REQUESTS_COLUMNS = [
    ("Request ID", "id"),
    ("Question", "question"),
    ("Related finding IDs", "related_finding_ids_text"),
    ("Priority", "priority"),
    ("Assigned recipient", "assigned_recipient"),
    ("Status", "status"),
    ("Management response", "management_response"),
    ("Reviewer follow-up", "reviewer_followup"),
    ("Created", "created_at"),
    ("Last updated", "updated_at"),
]


def _pdf_citation_text(citations: list[dict]) -> str:
    parts = []
    for c in citations or []:
        title = c.get("document_title") or "PDF"
        pages = f"p.{c['start_page']}" if c.get("end_page") in (None, c.get("start_page")) else f"p.{c['start_page']}-{c['end_page']}"
        parts.append(f"{title} {pages}")
    return "; ".join(parts)


def _excel_citation_text(citations: list[dict]) -> str:
    parts = []
    for c in citations or []:
        status = "verified" if c.get("exists") is True else "unverified" if c.get("exists") is False else "unchecked"
        parts.append(f"{c.get('workbook_label')}!{c.get('sheet')}!{c.get('ref')} [{c.get('kind')}] ({status})")
    return "; ".join(parts)


def findings_export_rows(findings: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Shared row shape for both the XLSX register and the printable
    package - built once so the two exports can never drift apart."""
    rows = []
    for f in findings:
        evidence_notes = f.get("evidence_notes") or ""
        pdf_evidence_text = f["pdf_evidence"] if f["origin"] == "ai" else evidence_notes
        workbook_evidence_text = _excel_citation_text(f.get("excel_citations", [])) or f["workbook_evidence"]
        pdf_citation_text = _pdf_citation_text(f.get("pdf_citations", []))
        if pdf_citation_text:
            pdf_evidence_text = f"{pdf_evidence_text} [{pdf_citation_text}]" if pdf_evidence_text else pdf_citation_text
        rows.append(
            {
                "id": f["id"],
                "origin_label": "AI-generated" if f["origin"] == "ai" else "Human-added",
                "title": f["title"],
                "classification": f["classification"],
                "ai_severity": f["severity"] if f["origin"] == "ai" else "",
                "adjusted_severity": f["adjusted_severity"] or "",
                "effective_severity": f["effective_severity"] or "",
                "review_status": f["review_status"],
                "resolution_status": f["resolution_status"],
                "assigned_owner": f["assigned_owner"],
                "duplicate_label": "Yes" if f["is_duplicate"] else "No",
                "duplicate_of": f["duplicate_of"] or "",
                "explanation": f["explanation"],
                "pdf_evidence_text": pdf_evidence_text,
                "workbook_evidence_text": workbook_evidence_text,
                "commercial_relevance": f["commercial_relevance"],
                "uncertainty": f["uncertainty"],
                "recommended_action": f["recommended_action"],
                "management_response": f["management_response"],
                "reviewer_notes": f["reviewer_notes"],
                "due_date_text": f["due_date_text"],
                "updated_at": f["updated_at"],
            }
        )
    return rows


def requests_export_rows(requests: list[workspaces.WorkspaceRequest]) -> list[dict[str, Any]]:
    return [
        {
            "id": r.id,
            "question": r.question,
            "related_finding_ids_text": ", ".join(r.related_finding_ids),
            "priority": r.priority,
            "assigned_recipient": r.assigned_recipient,
            "status": r.status,
            "management_response": r.management_response,
            "reviewer_followup": r.reviewer_followup,
            "created_at": r.created_at,
            "updated_at": r.updated_at,
        }
        for r in requests
    ]


def _safe_cell_text(value: str) -> str:
    """Prefix a leading formula-trigger character (=, +, -, @, tab, CR) with
    a literal quote so free text a reviewer typed can never be interpreted
    as a live spreadsheet formula when the export is opened elsewhere."""
    if value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return f"'{value}"
    return value


def _write_header_sheet(ws: Worksheet, project: store.Project, analysis: cross_format_analyses.CrossFormatAnalysis, workspace: workspaces.Workspace, export_kind: str) -> int:
    bold = Font(bold=True)
    now = datetime.now(timezone.utc).isoformat()
    lines = [
        ("Deal Intelligence Lab", export_kind),
        ("Project", f"{project.name} ({project.id})"),
        ("Analysis ID", analysis.id),
        ("Workspace ID", workspace.id),
        ("Exported at", now),
    ]
    row = 1
    for label, value in lines:
        ws.cell(row=row, column=1, value=label).font = bold
        ws.cell(row=row, column=2, value=_safe_cell_text(value))
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Disclaimer:").font = bold
    ws.cell(row=row + 1, column=1, value=_safe_cell_text(DISCLAIMER))
    ws.cell(row=row + 1, column=1).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=len(_FINDINGS_COLUMNS) or 2)
    return row + 3


def _write_table(ws: Worksheet, start_row: int, columns: list[tuple[str, str]], rows: list[dict[str, Any]]) -> None:
    bold = Font(bold=True)
    for col_index, (header, _) in enumerate(columns, start=1):
        cell = ws.cell(row=start_row, column=col_index, value=header)
        cell.font = bold
    for row_offset, row_data in enumerate(rows, start=1):
        for col_index, (_, key) in enumerate(columns, start=1):
            value = row_data.get(key, "")
            if isinstance(value, str):
                value = _safe_cell_text(value)
            ws.cell(row=start_row + row_offset, column=col_index, value=value)
    for col_index in range(1, len(columns) + 1):
        ws.column_dimensions[chr(64 + col_index) if col_index <= 26 else "A"].width = 24


def build_findings_workbook(
    project: store.Project,
    analysis: cross_format_analyses.CrossFormatAnalysis,
    workspace: workspaces.Workspace,
    findings: list[dict[str, Any]],
) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Findings register"
    start_row = _write_header_sheet(ws, project, analysis, workspace, "Findings register")
    _write_table(ws, start_row, _FINDINGS_COLUMNS, findings_export_rows(findings))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_requests_workbook(
    project: store.Project,
    analysis: cross_format_analyses.CrossFormatAnalysis,
    workspace: workspaces.Workspace,
    requests: list[workspaces.WorkspaceRequest],
) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Information requests"
    start_row = _write_header_sheet(ws, project, analysis, workspace, "Information-request list")
    _write_table(ws, start_row, _REQUESTS_COLUMNS, requests_export_rows(requests))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


_PRINT_CSS = """
<style>
  /* Always light, regardless of the viewing browser's theme: this page exists to be printed/saved as
     a PDF, and a dark-mode default would otherwise pair dark body text with a dark inherited
     background, making it unreadable on screen before anyone even gets to printing it. */
  :root { color-scheme: light; }
  html { background: #ffffff; }
  body { font-family: Georgia, 'Times New Roman', serif; color: #0f172a; background: #ffffff; max-width: 900px; margin: 0 auto; padding: 32px; line-height: 1.5; }
  h1 { font-size: 1.5rem; margin-bottom: 4px; }
  h2 { font-size: 1.1rem; margin-top: 28px; border-bottom: 1px solid #cbd5e1; padding-bottom: 4px; }
  .meta { color: #475569; font-size: 0.85rem; margin-bottom: 24px; }
  .meta div { margin-bottom: 2px; }
  .recommendation { display: inline-block; padding: 4px 10px; border: 1px solid #0f172a; border-radius: 4px; font-weight: bold; margin: 8px 0; }
  .body-text { white-space: pre-wrap; }
  table { border-collapse: collapse; width: 100%; margin-top: 8px; font-size: 0.78rem; }
  th, td { border: 1px solid #cbd5e1; padding: 6px 8px; text-align: left; vertical-align: top; }
  th { background: #e9eef5; }
  .disclaimer { margin-top: 32px; padding-top: 12px; border-top: 1px dashed #cbd5e1; font-size: 0.78rem; color: #475569; }
  @media print {
    a { color: inherit; text-decoration: none; }
    .no-print { display: none; }
  }
</style>
"""


def _memo_section(label: str, text: str) -> str:
    return f"<h2>{escape(label)}</h2><div class='body-text'>{escape(text) or '<em>Not yet completed.</em>'}</div>"


def _memo_body_html(memo: workspaces.WorkspaceMemo) -> str:
    status_label = "Approved" if memo.status == "approved" else "Draft"
    approval = (
        f"<div class='meta'>Approved by {escape(memo.approved_by or '')} at {escape(memo.approved_at or '')}</div>"
        if memo.status == "approved"
        else "<div class='meta'>Not yet approved - draft content, subject to change.</div>"
    )
    recommendation_label = memo.overall_recommendation.replace("_", " ")
    return (
        f"<div class='meta'><strong>Status: {status_label}</strong></div>"
        f"{approval}"
        f"<div class='recommendation'>Overall recommendation: {escape(recommendation_label)}</div>"
        + _memo_section("Executive conclusion", memo.executive_conclusion)
        + _memo_section("Transaction overview", memo.transaction_overview)
        + _memo_section("Critical issues", memo.critical_issues)
        + _memo_section("High-priority issues", memo.high_priority_issues)
        + _memo_section("Financial and valuation implications", memo.financial_valuation_implications)
        + _memo_section("Missing information", memo.missing_information)
        + _memo_section("Confirmed consistencies", memo.confirmed_consistencies)
        + _memo_section("Recommended next actions", memo.recommended_next_actions)
    )


def _identifying_header(project: store.Project, analysis: cross_format_analyses.CrossFormatAnalysis, workspace: workspaces.Workspace, title: str) -> str:
    now = datetime.now(timezone.utc).isoformat()
    return (
        f"<h1>{escape(title)}</h1>"
        "<div class='meta'>"
        f"<div>Project: {escape(project.name)} ({escape(project.id)})</div>"
        f"<div>Analysis ID: {escape(analysis.id)}</div>"
        f"<div>Workspace ID: {escape(workspace.id)}</div>"
        f"<div>Exported at: {escape(now)}</div>"
        "</div>"
    )


def build_memo_html(
    project: store.Project, analysis: cross_format_analyses.CrossFormatAnalysis, workspace: workspaces.Workspace, memo: workspaces.WorkspaceMemo
) -> str:
    return (
        "<!doctype html><html><head><meta charset='utf-8'>"
        f"<title>Executive Deal Memo - {escape(project.name)}</title>{_PRINT_CSS}</head><body>"
        + _identifying_header(project, analysis, workspace, "Executive Deal Memo")
        + _memo_body_html(memo)
        + f"<div class='disclaimer'>{escape(DISCLAIMER)}</div>"
        "</body></html>"
    )


def _findings_table_html(findings: list[dict[str, Any]]) -> str:
    columns = [c for c in _FINDINGS_COLUMNS if c[1] not in ("id",)]
    rows = findings_export_rows(findings)
    head = "".join(f"<th>{escape(h)}</th>" for h, _ in columns)
    body = ""
    for row in rows:
        cells = "".join(f"<td>{escape(str(row.get(key, '')))}</td>" for _, key in columns)
        body += f"<tr>{cells}</tr>"
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def _requests_table_html(requests: list[workspaces.WorkspaceRequest]) -> str:
    columns = [c for c in _REQUESTS_COLUMNS if c[1] != "id"]
    rows = requests_export_rows(requests)
    head = "".join(f"<th>{escape(h)}</th>" for h, _ in columns)
    body = ""
    for row in rows:
        cells = "".join(f"<td>{escape(str(row.get(key, '')))}</td>" for _, key in columns)
        body += f"<tr>{cells}</tr>"
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def build_package_html(
    project: store.Project,
    analysis: cross_format_analyses.CrossFormatAnalysis,
    workspace: workspaces.Workspace,
    findings: list[dict[str, Any]],
    requests: list[workspaces.WorkspaceRequest],
    memo: workspaces.WorkspaceMemo,
) -> str:
    return (
        "<!doctype html><html><head><meta charset='utf-8'>"
        f"<title>Decision Package - {escape(project.name)}</title>{_PRINT_CSS}</head><body>"
        + _identifying_header(project, analysis, workspace, "Decision-Ready Deal Package")
        + _memo_body_html(memo)
        + "<h2>Findings register</h2>"
        + _findings_table_html(findings)
        + "<h2>Information requests</h2>"
        + _requests_table_html(requests)
        + f"<div class='disclaimer'>{escape(DISCLAIMER)}</div>"
        "</body></html>"
    )
